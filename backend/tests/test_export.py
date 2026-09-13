"""Campaign results export tests: v2 columns (call status, duration,
avg e2e latency per call, one column per extracted field, flattened contact
custom_fields)."""

from __future__ import annotations

import csv
import io
from datetime import datetime
from typing import Any

from app.models import Call, Campaign, Contact, ExtractedField, Organization, Transcript
from app.services.export_service import export_csv
from conftest import auth_headers, register


def _seed(db: Any) -> Campaign:
    org = Organization(name="Acme", slug="acme")
    db.add(org)
    db.flush()
    campaign = Campaign(org_id=org.id, name="Absents - Aug")
    db.add(campaign)
    db.flush()
    contacted = Contact(
        campaign_id=campaign.id,
        org_id=org.id,
        name="Suresh",
        phone="+911234567890",
        status="completed",
        custom_fields={"student_name": "Aarav", "class_section": "10-B"},
    )
    pending = Contact(
        campaign_id=campaign.id,
        org_id=org.id,
        name="NoCall Yet",
        phone="+919876543210",
        status="queued",
        custom_fields={"student_name": "Ira"},
    )
    db.add_all([contacted, pending])
    db.flush()
    call = Call(
        campaign_id=campaign.id,
        contact_id=contacted.id,
        kind="phone",
        status="completed",
        duration_seconds=95.0,
    )
    db.add(call)
    db.flush()
    db.add_all(
        [
            Transcript(call_id=call.id, turn_index=0, speaker="agent", text="hi", e2e_ms=None),
            Transcript(call_id=call.id, turn_index=1, speaker="caller", text="fever", e2e_ms=None),
            Transcript(
                call_id=call.id, turn_index=1, speaker="agent", text="sorry to hear", e2e_ms=1000.0
            ),
            Transcript(
                call_id=call.id, turn_index=2, speaker="agent", text="ok", e2e_ms=2000.0
            ),
        ]
    )
    db.add(
        ExtractedField(
            call_id=call.id, field_name="reason_for_absence", field_value="fever"
        )
    )
    db.commit()
    return campaign


def test_export_has_v2_columns_and_values(session_factory):
    from sqlalchemy import select

    with session_factory() as db:
        campaign = _seed(db)
        csv_bytes = export_csv(db, campaign)

    rows = list(csv.reader(io.StringIO(csv_bytes.decode("utf-8-sig"))))
    header, data = rows[0], rows[1:]
    assert len(data) == 2

    # V2 columns present.
    for col in ("Contact Status", "Call Status", "Duration (s)", "E2E Latency (ms)"):
        assert col in header, f"missing {col}"
    assert "reason_for_absence" in header  # extracted field column
    assert "student_name" in header and "class_section" in header  # custom fields

    def cell(row_idx: int, col_name: str) -> str:
        return data[row_idx][header.index(col_name)]

    # Call-scoped values on the contacted row.
    assert cell(0, "Call Status") == "completed"
    assert float(cell(0, "Duration (s)")) == 95.0
    assert float(cell(0, "E2E Latency (ms)")) == 1500.0  # avg(1000, 2000); Nones skipped
    assert cell(0, "reason_for_absence") == "fever"
    assert cell(0, "student_name") == "Aarav"

    # Uncalled contact: empty call columns, custom fields still flattened.
    assert cell(1, "Call Status") == ""
    assert cell(1, "E2E Latency (ms)") == ""
    assert cell(1, "student_name") == "Ira"


# --- per-call export --------------------------------------------------------


def _seed_call_export(db: Any, org_id: int) -> int:
    campaign = Campaign(org_id=org_id, name="Absents - Aug")
    db.add(campaign)
    db.flush()
    contact = Contact(
        campaign_id=campaign.id,
        org_id=org_id,
        name="Suresh",
        phone="+911234567890",
        status="completed",
    )
    db.add(contact)
    db.flush()
    call = Call(
        campaign_id=campaign.id,
        contact_id=contact.id,
        org_id=org_id,
        kind="phone",
        status="completed",
        provider_call_id="plivo-123",
        started_at=datetime(2026, 9, 1, 10, 0, 0),
        ended_at=datetime(2026, 9, 1, 10, 1, 35),
        duration_seconds=95.0,
        outcome="resolved",
        flagged_for_human=True,
        cost={"total_usd": 0.02},
        latency={
            "stt_p50_ms": 300.0,
            "llm_p50_ms": 500.0,
            "tts_p50_ms": 400.0,
            "e2e_p50_ms": 1200.0,
        },
        summary="Student absent due to fever.",
        recording_url="https://example.test/rec.mp3",
    )
    db.add(call)
    db.flush()
    db.add_all(
        [
            Transcript(call_id=call.id, turn_index=0, speaker="agent", text="Hello."),
            Transcript(call_id=call.id, turn_index=1, speaker="caller", text="He has fever."),
            ExtractedField(call_id=call.id, field_name="reason_for_absence", field_value="fever"),
            ExtractedField(call_id=call.id, field_name="expected_return", field_value="aug 29"),
        ]
    )
    db.commit()
    return call.id


def test_call_export_csv_with_contact_fields_and_transcript(client, session_factory):
    token, user = register(client)
    with session_factory() as db:
        call_id = _seed_call_export(db, user["org_id"])

    resp = client.get(
        f"/api/calls/{call_id}/export",
        params={"format": "csv"},
        headers=auth_headers(token),
    )
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith("text/csv")
    assert f'filename="call-{call_id}.csv"' in resp.headers["content-disposition"]

    rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))
    header, row = rows[0], rows[1]
    assert len(rows) == 2, "single call export must contain exactly one data row"

    for col in (
        "Call ID",
        "Kind",
        "Status",
        "Provider Call ID",
        "Started",
        "Ended",
        "Duration (s)",
        "Contact Name",
        "Contact Phone",
        "Outcome",
        "Flagged For Human",
        "Cost",
        "Latency STT p50",
        "Latency LLM p50",
        "Latency TTS p50",
        "Latency E2E p50",
        "reason_for_absence",
        "expected_return",
        "Summary",
        "Recording URL",
        "Transcript",
    ):
        assert col in header, f"missing {col}"

    def cell(col_name: str) -> str:
        return row[header.index(col_name)]

    assert cell("Call ID") == str(call_id)
    assert cell("Kind") == "phone"
    assert cell("Status") == "completed"
    assert cell("Provider Call ID") == "plivo-123"
    assert cell("Started") == "2026-09-01T10:00:00"
    assert cell("Ended") == "2026-09-01T10:01:35"
    assert float(cell("Duration (s)")) == 95.0
    assert cell("Contact Name") == "Suresh"
    assert cell("Contact Phone") == "+911234567890"
    assert cell("Outcome") == "resolved"
    assert cell("Flagged For Human") == "True"
    assert cell("Cost") == "0.02"
    assert cell("Latency STT p50") == "300.0"
    assert cell("Latency LLM p50") == "500.0"
    assert cell("Latency TTS p50") == "400.0"
    assert cell("Latency E2E p50") == "1200.0"
    assert cell("reason_for_absence") == "fever"
    assert cell("expected_return") == "aug 29"
    assert cell("Summary") == "Student absent due to fever."
    assert cell("Recording URL") == "https://example.test/rec.mp3"
    assert "agent: Hello." in cell("Transcript")
    assert "caller: He has fever." in cell("Transcript")


def test_call_export_xlsx_contactless_playground(client, session_factory):
    token, user = register(client)
    with session_factory() as db:
        call = Call(
            kind="playground",
            org_id=user["org_id"],
            status="completed",
            duration_seconds=24.0,
        )
        db.add(call)
        db.flush()
        db.add(
            ExtractedField(call_id=call.id, field_name="reason_for_absence", field_value="fever")
        )
        db.commit()
        call_id = call.id

    resp = client.get(
        f"/api/calls/{call_id}/export",
        params={"format": "xlsx"},
        headers=auth_headers(token),
    )
    assert resp.status_code == 200, resp.text
    assert "spreadsheetml" in resp.headers["content-type"]
    assert f'filename="call-{call_id}.xlsx"' in resp.headers["content-disposition"]

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    header = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]

    assert "Call ID" in header
    assert header.index("Kind") < header.index("Contact Name")
    assert row[header.index("Kind")] == "playground"
    assert row[header.index("Status")] == "completed"
    assert float(row[header.index("Duration (s)")]) == 24.0
    assert not row[header.index("Contact Name")]
    assert not row[header.index("Contact Phone")]
    assert row[header.index("reason_for_absence")] == "fever"
    assert not row[header.index("Latency STT p50")]
    assert not row[header.index("Cost")]


def test_call_export_unknown_call_is_404(client):
    token, _ = register(client)
    resp = client.get(
        "/api/calls/999999/export",
        params={"format": "csv"},
        headers=auth_headers(token),
    )
    assert resp.status_code == 404


def test_call_export_foreign_org_call_is_404(client, session_factory):
    token_a, user_a = register(client, org_name="Org A", email="a@example.test")
    token_b, _ = register(client, org_name="Org B", email="b@example.test")
    with session_factory() as db:
        call_id = _seed_call_export(db, user_a["org_id"])

    resp = client.get(
        f"/api/calls/{call_id}/export",
        params={"format": "csv"},
        headers=auth_headers(token_b),
    )
    assert resp.status_code == 404


def test_call_export_invalid_format_is_422(client, session_factory):
    token, user = register(client)
    with session_factory() as db:
        call_id = _seed_call_export(db, user["org_id"])

    resp = client.get(
        f"/api/calls/{call_id}/export",
        params={"format": "pdf"},
        headers=auth_headers(token),
    )
    assert resp.status_code == 422
